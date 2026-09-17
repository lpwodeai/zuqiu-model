"""深度分析Excel文件中的赔率数据结构"""
import openpyxl
import os

def analyze_excel_odds(file_path, sheet_name):
    """分析单个工作表中的赔率数据"""
    wb = openpyxl.load_workbook(file_path, read_only=True)
    ws = wb[sheet_name]
    
    print(f"\n=== 工作表: {sheet_name} ===")
    
    # 读取所有行
    all_rows = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        all_rows.append(row)
    
    # 查找赔率相关的行
    odds_sections = []
    current_section = None
    
    for i, row in enumerate(all_rows):
        first_cell = str(row[0]) if row[0] else ''
        
        # 检测赔率区块
        if any(keyword in first_cell for keyword in ['胜平负', '让球', '总进球', '比分', '赔率', '盘口']):
            current_section = first_cell
            odds_sections.append({
                'section': first_cell,
                'start_row': i,
                'data': []
            })
        
        elif current_section and row[0] is not None and str(row[0]).strip() != '':
            if len(odds_sections) > 0:
                odds_sections[-1]['data'].append(row)
    
    # 打印赔率区块信息
    for section in odds_sections[:5]:
        print(f"\n区块: {section['section']}")
        print(f"起始行: {section['start_row']}")
        print(f"数据行数: {len(section['data'])}")
        if section['data']:
            for row in section['data'][:5]:
                display = [str(cell)[:20] if cell else '' for cell in row[:8]]
                print(f"  {display}")
    
    wb.close()
    return odds_sections

if __name__ == '__main__':
    import os
    base_dir = os.path.dirname(os.path.abspath(__file__))
    project_dir = os.path.dirname(base_dir)
    
    excel_path = os.path.join(project_dir, 'data', '2025-2026 英超 .xlsx')
    
    if os.path.exists(excel_path):
        # 分析前3个工作表
        wb = openpyxl.load_workbook(excel_path, read_only=True)
        sheet_names = wb.sheetnames[:3]
        wb.close()
        
        for sheet_name in sheet_names:
            analyze_excel_odds(excel_path, sheet_name)
    else:
        print(f"文件不存在: {excel_path}")